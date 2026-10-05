from openpyxl import load_workbook
wb = load_workbook('jobs.xlsx')
ws = wb.active

total = 0

for row in ws.iter_rows(min_row=2, values_only=True):
    print(row[0], row[1])
    total = total + int(row[1])
print("Total:", total)